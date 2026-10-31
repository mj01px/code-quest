"""Aplica na monografia do CodeQuest o resultado da auditoria de qualidade sistêmica.

Gera as versões marcada (antigo vermelho e tachado, novo vermelho e sublinhado)
e atualizada (antigo removido). Diferente de gerar_monografia_final.py, edita
também células de tabela (Quadro 41 e especificações dos RNFs), acrescenta
linhas de observação e um parágrafo novo em 4.1. O original não é alterado.

Uso: py atualizar_monografia_auditoria.py
"""

import copy
import sys
from pathlib import Path

from docx import Document
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

PASTA = Path(r"E:\Downloads_Archive\Monografia CodeQuest")
NOME = "[MONOGRAFIA] - CODEQUEST"
VERMELHO = RGBColor(255, 0, 0)

# Tabelas das especificações (doc.tables): Quadro 41 e uma por RNF.
Q41, RNF05, RNF06, RNF07, RNF12, RNF17 = 40, 45, 46, 47, 52, 57
RNF20, RNF21, RNF24, RNF26, RNF28, RNF30, RNF33 = 60, 61, 64, 66, 68, 70, 73

# troca: (ancora_inicio, ancora_fim ou None = até o fim, texto novo), com
# `celula` (tabela, linha, coluna) quando a âncora não é única no documento.
# fim: texto acrescentado ao fim do parágrafo que contém `paragrafo`.
# linha: nova linha (rótulo, texto) ao fim da tabela, no formato da última.
# depois: parágrafo novo logo após o que contém `paragrafo`.
EDICOES = [
    dict(
        tipo="fim",
        paragrafo="Os contadores de limitação de taxa e de cota ficam na memória",
        novo=(
            " Atrás do proxy do Next.js, o aluno é identificado pelo último endereço "
            "do cabeçalho X-Forwarded-For, acrescentado pelo proxy, e não pelo "
            "endereço da conexão, que seria o mesmo para todos os alunos (RNF05)."
        ),
    ),
    dict(
        tipo="troca",
        inicio="depois da aprovação do revisor,",
        fim="depois da aprovação do revisor,",
        novo="depois da aprovação de outro revisor (quem cria um conteúdo nunca o aprova),",
    ),
    dict(
        tipo="fim",
        paragrafo="Dos 37 requisitos, 36 estão implementados",
        novo=(
            " Das 104 regras de negócio associadas aos requisitos, 96 estão "
            "implementadas, cinco parcialmente (RN03 do RF06, RN02 do RF09, RN01 do "
            "RF11, RN01 do RF13 e RN01 do RF32), uma diverge do código (RN02 do RF15: "
            "o autor publica o próprio conteúdo aprovado sem a permissão de publicar) "
            "e duas pertencem ao RF37, planejado."
        ),
    ),
    dict(tipo="troca", celula=(Q41, 34, 3), inicio="Implementado", fim=None, novo="Parcial"),
    dict(
        tipo="troca",
        inicio="Dos 33 requisitos, 24 estão implementados",
        fim=None,
        novo=(
            "Dos 33 requisitos, 23 estão implementados, quatro deles com ressalva, "
            "oito parcialmente e dois planejados. As ressalvas tratam de situações "
            "que o critério original não previa: o consumo dos códigos de uso único "
            "do segundo fator sob pedidos simultâneos (RNF06), a página de entrada "
            "do painel administrativo padrão do Django (RNF07), a imutabilidade da "
            "auditoria garantida apenas pelo ORM (RNF12) e a conferência da cota do "
            "Judge0 entre processos (RNF17). Os parciais dependem, em geral, de "
            "infraestrutura ainda não adotada ou de verificação ainda não realizada. "
            "A limitação de taxa (RNF05) mantém os contadores na memória de cada "
            "processo e só será exata com um cache compartilhado, previsto na seção "
            "4.1, e a identificação do aluno atrás do proxy depende da configuração "
            "do servidor; o descarte da auditoria (RNF13) existe como comando, mas "
            "ainda sem agendamento; a verificação estrita de tipos (RNF26) cobre 4 "
            "dos 10 módulos; o contrato OpenAPI (RNF28) ainda deixa três rotas de "
            "fora; a acessibilidade (RNF30) e o suporte aos navegadores (RNF32) "
            "ainda não passaram por auditoria ou teste sistemático. A implantação em "
            "servidor próprio (RNF25) e a integração contínua (RNF27) estão planejadas."
        ),
    ),
    dict(
        tipo="troca",
        inicio="O cliente é identificado pelo endereço da conexão",
        fim="não pelo cabeçalho X-Forwarded-For.",
        novo=(
            "Atrás do proxy do Next.js, o cliente é identificado pelo último "
            "endereço do cabeçalho X-Forwarded-For (NUM_PROXIES igual a 1), "
            "acrescentado pelo proxy; isso exige que a API escute apenas no próprio "
            "servidor e que o Nginx sobrescreva esse cabeçalho com o endereço real "
            "da conexão."
        ),
    ),
    dict(
        tipo="fim",
        paragrafo="O que falta: Os contadores ficam na memória de cada processo.",
        novo=(
            " A identificação do cliente atrás do proxy depende da configuração do "
            "Nginx descrita no critério, prevista para a implantação (RNF25)."
        ),
    ),
    dict(
        tipo="linha",
        tabela=RNF06,
        rotulo="Observação",
        novo=(
            "Ressalva: o consumo dos códigos de uso único não é atômico, de modo que "
            "dois pedidos simultâneos com o mesmo código podem ser aceitos, e o "
            "último passo TOTP aceito não é guardado, o que permite repetir o mesmo "
            "código dentro da janela de tolerância. A correção consta na seção 4.1."
        ),
    ),
    dict(
        tipo="fim",
        paragrafo="Autenticação exigida por padrão em toda a API;",
        novo=(
            " Na aprovação de conteúdo vale o princípio dos quatro olhos: quem criou "
            "o conteúdo não pode aprová-lo, nem mesmo o administrador."
        ),
    ),
    dict(
        tipo="linha",
        tabela=RNF07,
        rotulo="Observação",
        novo=(
            "Ressalva: a rota do painel administrativo padrão do Django continua "
            "montada, e a página de entrada dele não passa pela limitação de taxa, "
            "pelo bloqueio por tentativas nem pelo segundo fator. A retirada dessa "
            "rota em produção consta na seção 4.1."
        ),
    ),
    dict(
        tipo="linha",
        tabela=RNF12,
        rotulo="Observação",
        novo=(
            "Ressalva: a imutabilidade é garantida pelo save() do ORM, sem gatilho no "
            "banco. Duas operações previstas passam por fora dele de propósito: a "
            "anonimização da conta (RNF14), que apaga e-mail, IP e navegador dos "
            "registros do titular por atualização em lote, e o descarte ao fim da "
            "retenção (RNF13), que remove os registros antigos."
        ),
    ),
    dict(
        tipo="linha",
        tabela=RNF17,
        rotulo="Observação",
        novo=(
            "Ressalva: a conferência da cota e a gravação da chamada não são "
            "atômicas entre processos, de modo que pedidos simultâneos em processos "
            "diferentes podem ultrapassar por pouco o limite diário."
        ),
    ),
    dict(
        tipo="troca",
        inicio="996 testes automatizados no backend",
        fim="total de 1.371 testes verdes.",
        novo=(
            "1.204 testes automatizados no backend (pytest), 488 no frontend (Jest), "
            "em 64 suítes, total de 1.692 testes."
        ),
    ),
    dict(
        tipo="fim",
        paragrafo="A cobertura é medida, mas não há percentual mínimo obrigatório",
        novo=(
            " Cinco testes de concorrência do backend verificam o bloqueio de linha "
            "e só passam no PostgreSQL, não no SQLite."
        ),
    ),
    dict(tipo="troca", inicio="40 arquivos de migration", fim="40 arquivos", novo="55 arquivos"),
    dict(tipo="troca", inicio="; Next 16.3.3 e React", fim="Next 16.3.3", novo="; Next 16.4.3"),
    dict(
        tipo="troca",
        inicio="44 arquivos de produção",
        fim="44 arquivos",
        novo="47 arquivos",
    ),
    dict(tipo="troca", celula=(RNF28, 5, 1), inicio="Implementado", fim=None, novo="Parcial"),
    dict(
        tipo="linha",
        tabela=RNF28,
        rotulo="Situação atual",
        novo=(
            "O que falta: três rotas ficam fora do contrato gerado, porque o "
            "drf-spectacular não consegue inferir o serializer delas: o início da "
            "desativação do 2FA, o detalhe de usuário do painel administrativo e o "
            "início de trilha. O padrão de erro vale para todas as rotas."
        ),
    ),
    dict(
        tipo="troca",
        inicio="(cerca de 200 ocorrências)",
        fim="(cerca de 200 ocorrências)",
        novo="(cerca de 390 ocorrências)",
    ),
    dict(
        tipo="troca",
        inicio="o contraste das cores e a navegação completa por teclado não foram medidos.",
        fim=None,
        novo=(
            "o contraste dos tons de texto secundário foi ajustado para ao menos "
            "4,5:1, mas a navegação completa por teclado não foi medida, e quatro "
            "janelas modais ainda não prendem o foco."
        ),
    ),
    dict(tipo="troca", inicio="Next.js 16.3.3 (App", fim="Next.js 16.3.3", novo="Next.js 16.4.3"),
    dict(
        tipo="depois",
        paragrafo="Abstração do motor de execução. Projeta-se",
        novo=(
            "Endurecimento da segurança e da qualidade. A auditoria de qualidade "
            "sistêmica, que cruzou os requisitos não funcionais e as regras de "
            "negócio com o código, deixou pendências registradas para as próximas "
            "versões. A limitação de taxa (RNF05) identifica o aluno pelo último "
            "endereço do cabeçalho X-Forwarded-For, acrescentado pelo proxy do "
            "Next.js; essa identificação só é confiável quando a API escuta apenas "
            "no próprio servidor e o Nginx sobrescreve o cabeçalho com o endereço "
            "real da conexão, condição que a implantação (RNF25) precisa garantir. "
            "Projetam-se ainda: tornar atômicos o consumo dos códigos de uso único "
            "do segundo fator e a conferência da cota diária do Judge0 entre "
            "processos, com bloqueio de linha no banco, e recusar a repetição de um "
            "código TOTP dentro da mesma janela (RNF06 e RNF17); proteger a trilha "
            "de auditoria também no banco, por gatilho (RNF12); retirar a rota do "
            "painel administrativo padrão do Django em produção (RNF07); incluir no "
            "contrato OpenAPI as três rotas que ainda ficam fora dele (RNF28); "
            "prender o foco do teclado em todas as janelas modais e criar telas de "
            "erro gerais para as rotas da aplicação (RNF30); estender a verificação "
            "estrita de tipos aos demais módulos (RNF26); e fazer a guarda de "
            "publicação conferir toda a hierarquia do conteúdo, e não apenas os "
            "filhos diretos."
        ),
    ),
]


def paragrafos(doc):
    """Parágrafos do corpo e das células de tabela, sem repetir célula mesclada."""
    yield from doc.paragraphs
    for tabela in doc.tables:
        vistos = set()
        for linha in tabela.rows:
            for celula in linha.cells:
                if id(celula._tc) in vistos:
                    continue
                vistos.add(id(celula._tc))
                yield from celula.paragraphs


def achar(doc, trecho, celula=None):
    if celula is not None:
        t, r, c = celula
        candidatos = doc.tables[t].rows[r].cells[c].paragraphs
    else:
        candidatos = list(paragrafos(doc))
    achados = [p for p in candidatos if trecho in p.text]
    if len(achados) != 1:
        sys.exit(f"Âncora {trecho[:50]!r}: {len(achados)} parágrafos (esperado 1).")
    return achados[0]


def dividir(par, run, pos):
    texto = run.text
    if pos <= 0 or pos >= len(texto):
        return
    novo = copy.deepcopy(run._r)
    run._r.addnext(novo)
    run.text = texto[:pos]
    Run(novo, par).text = texto[pos:]


def so_texto(run):
    return all(el.tag.endswith(("}rPr", "}t")) for el in run._r.iterchildren())


def runs_do_trecho(par, ini, fim):
    for ponto in (ini, fim):
        acc = 0
        for run in par.runs:
            tam = len(run.text)
            if acc < ponto < acc + tam:
                if not so_texto(run):
                    sys.exit("Run com tab/quebra na borda da edição; ajuste a âncora.")
                dividir(par, run, ponto - acc)
                break
            acc += tam
    escolhidos, acc = [], 0
    for run in par.runs:
        tam = len(run.text)
        if tam and acc >= ini and acc + tam <= fim:
            escolhidos.append(run)
        acc += tam
    return escolhidos


def marcar_novo(run, marcada):
    run.font.strike = None
    if marcada:
        run.font.color.rgb = VERMELHO
        run.font.underline = True


def run_apos(par, modelo, depois, texto):
    elem = copy.deepcopy(modelo._r)
    depois._r.addnext(elem)
    run = Run(elem, par)
    run.text = texto
    return run


def so_um_run(par, texto, marcada):
    """Deixa o parágrafo com o primeiro run (formatação) e o texto dado."""
    runs = par.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    runs[0].text = texto
    marcar_novo(runs[0], marcada)


def aplicar(doc, ed, marcada):
    tipo = ed["tipo"]
    if tipo == "fim":
        par = achar(doc, ed["paragrafo"])
        marcar_novo(run_apos(par, par.runs[-1], par.runs[-1], ed["novo"]), marcada)
    elif tipo == "troca":
        par = achar(doc, ed["inicio"], ed.get("celula"))
        texto = par.text
        ini = texto.index(ed["inicio"])
        fim = len(texto) if ed["fim"] is None else texto.index(ed["fim"], ini) + len(ed["fim"])
        antigos = runs_do_trecho(par, ini, fim)
        if not antigos:
            sys.exit(f"Troca {ed['inicio'][:40]!r}: nenhum run no trecho.")
        marcar_novo(run_apos(par, antigos[0], antigos[-1], ed["novo"]), marcada)
        for r in antigos:
            if marcada:
                r.font.color.rgb = VERMELHO
                r.font.strike = True
            else:
                r._r.getparent().remove(r._r)
    elif tipo == "linha":
        tabela = doc.tables[ed["tabela"]]
        modelo = tabela.rows[-1]._tr
        nova = copy.deepcopy(modelo)
        modelo.addnext(nova)
        linha = tabela.rows[-1]
        so_um_run(linha.cells[0].paragraphs[0], ed["rotulo"], marcada)
        so_um_run(linha.cells[1].paragraphs[0], ed["novo"], marcada)
        for celula in linha.cells:
            for p in celula.paragraphs[1:]:
                p._p.getparent().remove(p._p)
    elif tipo == "depois":
        par = achar(doc, ed["paragrafo"])
        elem = copy.deepcopy(par._p)
        par._p.addnext(elem)
        so_um_run(Paragraph(elem, par._parent), ed["novo"], marcada)


def gerar(origem, destino, marcada):
    doc = Document(origem)
    for ed in EDICOES:
        aplicar(doc, ed, marcada)
    doc.save(destino)
    print("gerado:", destino)


if __name__ == "__main__":
    origem = PASTA / f"{NOME}.docx"
    gerar(origem, PASTA / f"{NOME} - marcada.docx", marcada=True)
    gerar(origem, PASTA / f"{NOME} - atualizada.docx", marcada=False)
